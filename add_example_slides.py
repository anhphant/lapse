"""Thêm 3 slide ví dụ (pygame / tkinter / chat mode) vào timelapse_slides.pptx."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

SRC = "timelapse_slides.pptx"

DARK = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x40, 0x40, 0x40)
CODE = RGBColor(0x1A, 0x1A, 0x1A)


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.name = "Calibri"
    p.font.color.rgb = DARK


def add_code(slide, lines, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(14)
        p.font.name = "Consolas"
        p.font.color.rgb = CODE
        p.space_after = Pt(2)


def add_bullets(slide, items, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(16)
        p.font.name = "Calibri"
        p.font.color.rgb = GREY
        p.space_after = Pt(12)


prs = Presentation(SRC)
blank = prs.slide_layouts[6]

# ---------- pygame ----------
s = prs.slides.add_slide(blank)
add_title(s, "Ví dụ: pygame")
add_code(s, [
    "import pygame",
    "pygame.init()",
    "screen = pygame.display.set_mode((400, 300))",
    "running = True",
    "while running:",
    "    for e in pygame.event.get():",
    "        if e.type == pygame.QUIT:",
    "            running = False",
    "    screen.fill((0, 0, 0))",
    "    pygame.draw.circle(screen, (0, 255, 0), (200, 150), 40)",
    "    pygame.display.flip()",
    "pygame.quit()",
], Inches(0.5), Inches(1.4), Inches(6.8), Inches(5.2))
add_bullets(s, [
    "Vòng lặp game chuẩn: init → xử lý sự kiện → vẽ → flip → quit.",
    "Dùng đúng API pygame (sự kiện, vẽ, cập nhật màn hình).",
], Inches(7.7), Inches(1.6), Inches(5.1), Inches(4.5))

# ---------- tkinter ----------
s = prs.slides.add_slide(blank)
add_title(s, "Ví dụ: tkinter")
add_code(s, [
    "import tkinter as tk",
    "root = tk.Tk()",
    "root.title(\"Demo\")",
    "tk.Label(root, text=\"Xin chào\").pack()",
    "tk.Button(root, text=\"OK\", command=root.destroy).pack()",
    "root.mainloop()",
], Inches(0.5), Inches(1.4), Inches(6.8), Inches(3.2))
add_bullets(s, [
    "Dựng cửa sổ desktop với widget chuẩn (Label, Button).",
    "Chạy vòng lặp sự kiện mainloop() để nhận tương tác.",
], Inches(7.7), Inches(1.6), Inches(5.1), Inches(4.5))

# ---------- chat mode ----------
s = prs.slides.add_slide(blank)
add_title(s, "Ví dụ: chat mode")
add_code(s, [
    "Người dùng: viết script đếm số dòng trong file",
    "AI: trả về code kèm giải thích",
    "",
    "Người dùng: bỏ dòng trống, đếm cả thư mục con",
    "AI: cập nhật code theo yêu cầu mới",
    "",
    "Người dùng: thêm thanh tiến trình",
    "AI: chỉnh tiếp, giữ đúng ngữ cảnh",
], Inches(0.5), Inches(1.4), Inches(6.8), Inches(4.6))
add_bullets(s, [
    "Hội thoại nhiều lượt, AI giữ ngữ cảnh.",
    "Chỉnh sửa dần theo phản hồi của người dùng.",
], Inches(7.7), Inches(1.6), Inches(5.1), Inches(4.5))

prs.save(SRC)
print(f"OK -> {SRC} | tong slide: {len(prs.slides.__iter__.__self__._sldIdLst)}")
