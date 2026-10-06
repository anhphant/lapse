"""Tạo slide .pptx trình bày kết quả time-lapse.
Tự chèn figure3.1.png nếu có trong thư mục; nếu chưa có thì vẽ ô placeholder.
Chạy lại sau khi copy figure3.1.png vào đây để chèn ảnh thật.
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = "timelapse_slides.pptx"
FIG = "figure3.1.png"

DARK = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x40, 0x40, 0x40)
ACCENT = RGBColor(0x2E, 0x74, 0xB5)


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(0.9))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.name = "Calibri"
    p.font.color.rgb = DARK


def add_bullets(slide, items, left, top, width, height, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.name = "Calibri"
        p.font.color.rgb = GREY
        p.space_after = Pt(10)
    return box


def add_placeholder(slide, left, top, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xEF, 0xEF, 0xEF)
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1.5)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "figure3.1.png\n\n(chép file vào thư mục rồi chạy lại make_slides.py)"
    p.font.size = Pt(14)
    p.font.name = "Calibri"
    p.font.color.rgb = GREY
    p.alignment = PP_ALIGN.CENTER
    return shape


def add_code(slide, lines, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(13)
        p.font.name = "Consolas"
        p.font.color.rgb = RGBColor(0x0B, 0x3D, 0x2E) if line.strip().startswith("#") else RGBColor(0x1A, 0x1A, 0x1A)
        p.space_after = Pt(2)
    return box


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# ---------- Slide 1: Tuân thủ chỉ dẫn tốt hơn ----------
s1 = prs.slides.add_slide(blank)
add_title(s1, "Tuân thủ chỉ dẫn tốt hơn")
add_bullets(s1, [
    "Thư viện đúng yêu cầu: opencv-python (webcam + xử lý khung hình) + mss (chụp màn hình tốc độ cao).",
    "Cửa sổ Preview thời gian thực gộp màn hình + webcam: chế độ side-by-side hoặc picture-in-picture.",
    "Cơ chế time-lapse chuẩn: INTERVAL cấu hình (1 khung / N giây), ghi trực tiếp ra file .mp4.",
    "Thông số ở đầu file (INTERVAL, WEBCAM_ID, FPS, Resolution), thoát an toàn bằng phím 'q' và giải phóng tài nguyên.",
], Inches(0.5), Inches(1.4), Inches(7.0), Inches(5.6), size=17)

img_left, img_top = Inches(7.8), Inches(1.5)
if os.path.exists(FIG):
    s1.shapes.add_picture(FIG, img_left, img_top, width=Inches(5.0))
    cap = s1.shapes.add_textbox(img_left, Inches(5.9), Inches(5.0), Inches(0.5))
    cap.text_frame.text = "Hình 3.1: Kết quả chạy script time-lapse"
    cap.text_frame.paragraphs[0].font.size = Pt(12)
    cap.text_frame.paragraphs[0].font.italic = True
    cap.text_frame.paragraphs[0].font.color.rgb = GREY
else:
    add_placeholder(s1, img_left, img_top, Inches(5.0), Inches(3.6))

# ---------- Slide 2: Cải thiện khả năng dùng thư viện ngoài ----------
s2 = prs.slides.add_slide(blank)
add_title(s2, "Cải thiện khả năng dùng thư viện ngoài")
add_bullets(s2, [
    "Chọn đúng thư viện theo mục đích: mss chụp màn hình ở mức hệ điều hành (nhanh, mượt) thay cho pyautogui chậm hơn.",
    "Dùng đúng API theo đặc tả từng thư viện: đổi kênh màu (BGRA → BGR), đặt codec (mp4v) và FPS cho VideoWriter.",
    "Kết hợp thư viện bằng numpy để gộp/scale khung hình (hstack, resize) trước khi ghi video.",
], Inches(0.5), Inches(1.4), Inches(6.3), Inches(5.4), size=16)

add_code(s2, [
    "# mss: chụp màn hình nhanh (trả về BGRA)",
    "with mss.mss() as sct:",
    "    shot = sct.grab(mon)",
    "    screen = cv2.cvtColor(np.array(shot), cv2.COLOR_BGRA2BGR)",
    "",
    "# opencv: gộp webcam + màn hình rồi ghi .mp4",
    "frame = np.hstack([screen, webcam])",
    "writer = cv2.VideoWriter('out.mp4',",
    "    cv2.VideoWriter_fourcc(*'mp4v'), 10, (w, h))",
    "writer.write(frame)",
], Inches(7.2), Inches(1.5), Inches(5.6), Inches(5.2))

prs.save(OUT)
status = "da chen" if os.path.exists(FIG) else "chua co - placeholder"
print(f"OK -> {OUT}  (figure3.1.png: {status})")
